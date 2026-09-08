import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

template_path = r'C:\Users\misha\OneDrive\Desktop\PHEX-HEV-EMS-main\Internshi PPT_Template _ BTech.pptx'
output_path = r'C:\Users\misha\OneDrive\Desktop\PHEX-HEV-EMS-main\Final_Internship_Viva_Presentation_Mishael_Julian.pptx'

prs = Presentation(template_path)
print(f'Loaded template with {len(prs.slides)} slides.')
