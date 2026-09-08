import pptx

prs = pptx.Presentation(r'C:\Users\misha\OneDrive\Desktop\PHEX-HEV-EMS-main\Internshi PPT_Template _ BTech.pptx')
for idx, slide in enumerate(prs.slides):
    print(f'=== SLIDE {idx+1} ===')
    for s in slide.shapes:
        print(f'  Shape: {s.name} (type: {s.shape_type})')
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                print(f'    P: "{p.text}" (font: {p.font.name}, size: {p.font.size})')
        elif s.has_table:
            print(f'    Table: {len(s.table.rows)} rows x {len(s.table.columns)} cols')
            for r in s.table.rows:
                print('      ' + ' | '.join([c.text.strip().replace('\n', ' ') for c in r.cells]))
